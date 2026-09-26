"""Exercise real Tk controls against an isolated database, including geometry."""
import os,sys,tempfile
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'todo'))
temp=tempfile.TemporaryDirectory()
os.environ['TODO_DB_PATH']=str(Path(temp.name)/'ui.db')
from app import ToDoApp
from database import connect
from screens.products import ProductEditor, ProductsFrame
from screens.payment import PaymentDialog
from services.cash import open_session
from services.receipts import export_receipt_pdf

def fail(title,message,**kwargs):
    raise AssertionError(str(title)+': '+str(message))
def descendants(w):
    for child in w.winfo_children():
        yield child
        yield from descendants(child)
def button(w, label):
    return next(x for x in descendants(w) if x.winfo_class() in ('TButton','Button') and label in str(x.cget('text')))
def visible(w):
    app.update()
    assert w.winfo_viewable(), str(w)
    top=w.winfo_toplevel()
    assert w.winfo_rooty()+w.winfo_height() <= top.winfo_rooty()+top.winfo_height(), str(w.cget('text'))
    assert w.winfo_rootx()+w.winfo_width() <= top.winfo_rootx()+top.winfo_width(), str(w.cget('text'))

with patch('tkinter.messagebox.showerror',fail), patch('tkinter.messagebox.showwarning',fail), patch('tkinter.messagebox.showinfo',lambda *a,**k:None):
    app=ToDoApp()
    app.report_callback_exception=lambda t,v,tb: fail('Tk callback',v)
    app.login_pin.set('1234');app.login()
    app.geometry('1100x640');app.update()
    done=__import__('tkinter').BooleanVar(value=False)
    app.after(200,lambda:done.set(True));app.wait_variable(done)
    app.show('stock');app.update()
    button(app.current,'Articles').invoke()
    app.update()
    editor=ProductEditor(app)
    editor.name.set('TEST - Rice')
    with patch('screens.products.simpledialog.askstring',return_value='Test groceries'):
        editor.add_category()
    editor.buy.set('2');editor.sell.set('3')
    editor.bar.set('TEST123')
    button(editor,'Enregistrer').invoke();app.update()
    with connect() as c:
        pid=c.execute("SELECT id FROM products WHERE name='TEST - Rice'").fetchone()[0]
        assert c.execute('SELECT category_id FROM products WHERE id=?',(pid,)).fetchone()[0]
    app.show('purchases');app.update()
    purchase=app.current
    from screens.suppliers import SupplierEditor, SuppliersFrame
    button(purchase,'+').invoke();app.update()
    supplier_editor=next(w for w in descendants(app) if isinstance(w,SupplierEditor))
    supplier_editor.name.set('TEST Supplier');supplier_editor.phone.set('0600000000')
    button(supplier_editor,'Enregistrer').invoke();app.update()
    assert purchase.supplier_id is not None
    supplier_id=purchase.supplier_id
    purchase.invoice.set('TEST-001')
    purchase.prod.selection_set(purchase.prod.get_children()[0])
    with patch('tkinter.simpledialog.askfloat',side_effect=[5,2]):
        purchase.addline()
    confirm=button(purchase,'VALIDER');visible(confirm);confirm.invoke();app.update()
    with connect() as c:
        assert c.execute('SELECT stock_qty FROM products WHERE id=?',(pid,)).fetchone()[0]==5
        assert c.execute('SELECT supplier_id FROM purchases').fetchone()[0]==supplier_id
    app.show('management');app.update()
    supplier_cards=[w for w in descendants(app.current) if w.winfo_class()=='Label' and w.cget('text')=='Fournisseurs']
    assert len(supplier_cards)==1
    supplier_cards[0].event_generate('<Button-1>');app.update()
    assert isinstance(app.current,SuppliersFrame)
    app.current.tree.selection_set(str(supplier_id))
    button(app.current,'Modifier').invoke();app.update()
    supplier_editor=next(w for w in descendants(app) if isinstance(w,SupplierEditor))
    supplier_editor.name.set('TEST Supplier updated');button(supplier_editor,'Enregistrer').invoke()
    button(app.current,'Réceptions du fournisseur').invoke();app.update()
    history=next(w for w in descendants(app) if w.winfo_class()=='Toplevel')
    history_tree=next(w for w in descendants(history) if w.winfo_class()=='Treeview')
    assert len(history_tree.get_children())==1
    assert 'TEST-001' in history_tree.item(history_tree.get_children()[0],'values')
    history.destroy()
    open_session(app.user['id'],10000)
    app.show('sale');app.update()
    sale=app.current
    sale.entry.focus_set();app.update()
    button(app,'Clavier').invoke();app.update()
    from screens.virtual_keyboard import VirtualKeyboard
    assert VirtualKeyboard._instance and VirtualKeyboard._instance.winfo_exists()
    VirtualKeyboard._instance._close()
    assert str(pid) in sale.products.get_children(), 'Products without images must be searchable'
    # Photo-only products saved through the real editor receive a stable TODO-* internal barcode.
    photo_editor=ProductEditor(app)
    photo_editor.name.set('TEST Photo Internal')
    photo_editor.sell.set('4')
    photo_editor.img_rel='missing-test-image.jpg'
    button(photo_editor,'Enregistrer').invoke();app.update()
    with connect() as c:
        photo_internal_pid=c.execute("SELECT id FROM products WHERE name='TEST Photo Internal'").fetchone()[0]
        internal_barcode=c.execute("SELECT barcode FROM product_barcodes WHERE product_id=?",(photo_internal_pid,)).fetchone()[0]
        assert internal_barcode==f'TODO-{photo_internal_pid:08d}'
        c.execute("INSERT INTO products(name,sale_price_cents,active,image_path) VALUES('TEST Photo Regular',500,1,'missing-test-image-2.jpg')")
        photo_regular_pid=c.execute("SELECT last_insert_rowid()").fetchone()[0]
        c.execute("INSERT INTO product_barcodes(product_id,barcode,qty_multiplier) VALUES(?,?,1)",(photo_regular_pid,'REGULAR123'))
        c.execute("INSERT INTO products(name,sale_price_cents,active) VALUES('TEST Inactive Export',999,0)")
    # Shop-to-shop catalogue export must never leak device-local TODO-* barcodes.
    export_window=__import__('tkinter').Toplevel(app)
    products_frame=ProductsFrame(export_window);products_frame.pack(fill='both',expand=True)
    export_path=Path(temp.name)/'catalogue-exchange.xlsx'
    with patch('screens.products.filedialog.asksaveasfilename',return_value=str(export_path)):
        products_frame.export_catalogue()
    from openpyxl import load_workbook
    exported_wb=load_workbook(export_path,read_only=True,data_only=True)
    exported=list(exported_wb.active.iter_rows(values_only=True))
    exported_wb.close()
    header=exported[0];barcode_col=header.index('barcode');name_col=header.index('article')
    rows_by_name={}
    for row in exported[1:]:rows_by_name.setdefault(row[name_col],[]).append(row)
    assert rows_by_name['TEST Photo Internal'][0][barcode_col] in (None,''), 'TODO-* internal barcode must not be exported'
    assert any(row[barcode_col]=='REGULAR123' for row in rows_by_name['TEST Photo Regular']), 'Real barcode must remain in catalogue export'
    assert 'TEST Inactive Export' not in rows_by_name, 'Inactive products must stay out of shop-to-shop catalogue export'
    export_window.destroy()
    # Inactive local barcodes must not block shop-to-shop import as conflicts.
    from openpyxl import Workbook
    import_path=Path(temp.name)/'catalogue-import-inactive-conflict.xlsx'
    wb=Workbook();ws=wb.active
    ws.append(["product key","barcode","article","famille","prix achat","prix vente","stock","alerte","barcode label","multiplicateur","prix pack","sku","fraction"])
    ws.append(["TODO-000001","INACTIVE123","Imported Active Product","Général","1.00","2.00",3,0,"",1,None,"",0]);wb.save(import_path);wb.close()
    with connect() as c:
        c.execute("INSERT INTO products(name,sale_price_cents,active) VALUES('Old Inactive Local',100,0)")
        inactive_pid=c.execute("SELECT last_insert_rowid()").fetchone()[0]
        c.execute("INSERT INTO product_barcodes(product_id,barcode,qty_multiplier) VALUES(?,?,1)",(inactive_pid,'INACTIVE123'))
    import_host=__import__('tkinter').Toplevel(app)
    products_import=ProductsFrame(import_host);products_import.pack(fill='both',expand=True);app.update()
    def accept_import_preview():
        preview=next(w for w in products_import.winfo_children() if w.winfo_class()=='Toplevel')
        button(preview,'Importer 1').invoke()
    app.after(150,accept_import_preview)
    with patch('screens.products.filedialog.askopenfilename',return_value=str(import_path)), \
         patch('screens.products.messagebox.askyesnocancel',side_effect=AssertionError('Inactive barcode must not trigger conflict dialog')):
        products_import.import_excel()
    app.update()
    with connect() as c:
        assert c.execute("SELECT 1 FROM products p JOIN product_barcodes b ON b.product_id=p.id WHERE p.active=1 AND p.name='Imported Active Product' AND b.barcode='INACTIVE123'").fetchone()
    import_host.destroy()
    # Legacy catalogues may contain device-local TODO-* pseudo barcodes. Import the
    # product but never persist that internal id as a real barcode on this device.
    legacy_path=Path(temp.name)/'catalogue-legacy-internal.xlsx'
    wb=Workbook();ws=wb.active
    ws.append(["product key","barcode","article","famille","prix achat","prix vente","stock","alerte","barcode label","multiplicateur","prix pack","sku","fraction"])
    ws.append(["TODO-LEGACY","TODO-00000042","Imported Legacy Internal","Général","1.00","2.00",1,0,"",1,None,"",0]);wb.save(legacy_path);wb.close()
    legacy_host=__import__('tkinter').Toplevel(app)
    legacy_import=ProductsFrame(legacy_host);legacy_import.pack(fill='both',expand=True);app.update()
    def accept_legacy_preview():
        preview=next(w for w in legacy_import.winfo_children() if w.winfo_class()=='Toplevel')
        button(preview,'Importer 1').invoke()
    app.after(150,accept_legacy_preview)
    with patch('screens.products.filedialog.askopenfilename',return_value=str(legacy_path)):
        legacy_import.import_excel()
    with connect() as c:
        legacy_pid=c.execute("SELECT id FROM products WHERE active=1 AND name='Imported Legacy Internal'").fetchone()[0]
        assert c.execute("SELECT COUNT(*) FROM product_barcodes WHERE product_id=?",(legacy_pid,)).fetchone()[0]==0
    legacy_host.destroy()
    # Multi-category catalogue exchange must preserve every family, not only the primary one.
    with connect() as c:
        c.execute("INSERT OR IGNORE INTO categories(name) VALUES('Family A')")
        family_a=c.execute("SELECT id FROM categories WHERE name='Family A'").fetchone()[0]
        c.execute("INSERT OR IGNORE INTO categories(name) VALUES('Family B')")
        family_b=c.execute("SELECT id FROM categories WHERE name='Family B'").fetchone()[0]
        c.execute("INSERT INTO products(name,sale_price_cents,active,category_id) VALUES('TEST Multi Family',700,1,?)",(family_a,))
        multi_pid=c.execute("SELECT last_insert_rowid()").fetchone()[0]
        c.executemany("INSERT OR IGNORE INTO product_categories(product_id,category_id) VALUES(?,?)",[(multi_pid,family_a),(multi_pid,family_b)])
        c.execute("INSERT INTO product_barcodes(product_id,barcode,qty_multiplier) VALUES(?,?,1)",(multi_pid,'MULTIFAM123'))
    multi_export=Path(temp.name)/'catalogue-multi-family.xlsx'
    multi_export_host=__import__('tkinter').Toplevel(app)
    multi_products=ProductsFrame(multi_export_host);multi_products.pack(fill='both',expand=True);app.update()
    with patch('screens.products.filedialog.asksaveasfilename',return_value=str(multi_export)):
        multi_products.export_catalogue()
    multi_export_host.destroy()
    # Keep only this product in the round-trip file so unrelated existing products
    # cannot open barcode-conflict dialogs and block the acceptance test.
    multi_wb=load_workbook(multi_export)
    multi_ws=multi_wb.active
    multi_headers=[cell.value for cell in multi_ws[1]]
    name_col=multi_headers.index('article')+1
    for row_idx in range(multi_ws.max_row,1,-1):
        if multi_ws.cell(row_idx,name_col).value!='TEST Multi Family':
            multi_ws.delete_rows(row_idx,1)
    multi_wb.save(multi_export);multi_wb.close()
    # Remove the source product/categories so the import proves round-trip preservation.
    with connect() as c:
        c.execute("DELETE FROM products WHERE id=?",(multi_pid,))
        c.execute("DELETE FROM categories WHERE id IN (?,?)",(family_a,family_b))
    multi_import_host=__import__('tkinter').Toplevel(app)
    multi_import=ProductsFrame(multi_import_host);multi_import.pack(fill='both',expand=True);app.update()
    def accept_multi_preview():
        preview=next(w for w in multi_import.winfo_children() if w.winfo_class()=='Toplevel')
        button(preview,'Importer').invoke()
    app.after(150,accept_multi_preview)
    with patch('screens.products.filedialog.askopenfilename',return_value=str(multi_export)):
        multi_import.import_excel()
    with connect() as c:
        imported_multi=c.execute("SELECT id FROM products WHERE active=1 AND name='TEST Multi Family'").fetchone()[0]
        families={r[0] for r in c.execute("""SELECT cat.name FROM product_categories pc JOIN categories cat ON cat.id=pc.category_id WHERE pc.product_id=?""",(imported_multi,)).fetchall()}
        assert families=={'Family A','Family B'}, families
    multi_import_host.destroy()
    # Quantity-pricing offers must survive catalogue exchange exactly.
    with connect() as c:
        c.execute("INSERT INTO products(name,sale_price_cents,active) VALUES('TEST Quantity Offers',1000,1)")
        offers_pid=c.execute("SELECT last_insert_rowid()").fetchone()[0]
        c.execute("INSERT INTO product_barcodes(product_id,barcode,qty_multiplier) VALUES(?,?,1)",(offers_pid,'OFFERS123'))
        c.executemany("INSERT INTO quantity_prices(product_id,min_qty,unit_price_cents,pricing_mode) VALUES(?,?,?,?)",[
            (offers_pid,2,800,'UNIT'),(offers_pid,3,2500,'BUNDLE')
        ])
    offers_export=Path(temp.name)/'catalogue-quantity-offers.xlsx'
    offers_export_host=__import__('tkinter').Toplevel(app)
    offers_products=ProductsFrame(offers_export_host);offers_products.pack(fill='both',expand=True);app.update()
    with patch('screens.products.filedialog.asksaveasfilename',return_value=str(offers_export)):
        offers_products.export_catalogue()
    offers_export_host.destroy()
    offers_wb=load_workbook(offers_export)
    offers_ws=offers_wb.active
    offers_headers=[cell.value for cell in offers_ws[1]]
    offers_name_col=offers_headers.index('article')+1
    for row_idx in range(offers_ws.max_row,1,-1):
        if offers_ws.cell(row_idx,offers_name_col).value!='TEST Quantity Offers':
            offers_ws.delete_rows(row_idx,1)
    offers_wb.save(offers_export);offers_wb.close()
    with connect() as c:
        c.execute("DELETE FROM products WHERE id=?",(offers_pid,))
    offers_import_host=__import__('tkinter').Toplevel(app)
    offers_import=ProductsFrame(offers_import_host);offers_import.pack(fill='both',expand=True);app.update()
    def accept_offers_preview():
        preview=next(w for w in offers_import.winfo_children() if w.winfo_class()=='Toplevel')
        button(preview,'Importer 1').invoke()
    app.after(150,accept_offers_preview)
    with patch('screens.products.filedialog.askopenfilename',return_value=str(offers_export)):
        offers_import.import_excel()
    with connect() as c:
        imported_offers_pid=c.execute("SELECT id FROM products WHERE active=1 AND name='TEST Quantity Offers'").fetchone()[0]
        imported_offers=[tuple(r) for r in c.execute("SELECT min_qty,unit_price_cents,pricing_mode FROM quantity_prices WHERE product_id=? ORDER BY min_qty",(imported_offers_pid,)).fetchall()]
        assert imported_offers==[(2.0,800,'UNIT'),(3.0,2500,'BUNDLE')], imported_offers
    offers_import_host.destroy()
    # Search metadata (alias + supplier code) must survive catalogue exchange.
    with connect() as c:
        c.execute("INSERT INTO products(name,sale_price_cents,active,sku,alias,supplier_code) VALUES('TEST Search Metadata',500,1,'META-SKU','Alias Unique','SUP-XYZ')")
        meta_pid=c.execute("SELECT last_insert_rowid()").fetchone()[0]
        c.execute("INSERT INTO product_barcodes(product_id,barcode,qty_multiplier) VALUES(?,?,1)",(meta_pid,'META123'))
    meta_export=Path(temp.name)/'catalogue-search-metadata.xlsx'
    meta_export_host=__import__('tkinter').Toplevel(app)
    meta_products=ProductsFrame(meta_export_host);meta_products.pack(fill='both',expand=True);app.update()
    with patch('screens.products.filedialog.asksaveasfilename',return_value=str(meta_export)):
        meta_products.export_catalogue()
    meta_export_host.destroy()
    meta_wb=load_workbook(meta_export)
    meta_ws=meta_wb.active
    meta_headers=[cell.value for cell in meta_ws[1]]
    meta_name_col=meta_headers.index('article')+1
    for row_idx in range(meta_ws.max_row,1,-1):
        if meta_ws.cell(row_idx,meta_name_col).value!='TEST Search Metadata':
            meta_ws.delete_rows(row_idx,1)
    meta_wb.save(meta_export);meta_wb.close()
    with connect() as c:
        c.execute("DELETE FROM products WHERE id=?",(meta_pid,))
    meta_import_host=__import__('tkinter').Toplevel(app)
    meta_import=ProductsFrame(meta_import_host);meta_import.pack(fill='both',expand=True);app.update()
    def accept_meta_preview():
        preview=next(w for w in meta_import.winfo_children() if w.winfo_class()=='Toplevel')
        button(preview,'Importer 1').invoke()
    app.after(150,accept_meta_preview)
    with patch('screens.products.filedialog.askopenfilename',return_value=str(meta_export)):
        meta_import.import_excel()
    with connect() as c:
        meta_row=c.execute("SELECT sku,alias,supplier_code FROM products WHERE active=1 AND name='TEST Search Metadata'").fetchone()
        assert tuple(meta_row)==('META-SKU','Alias Unique','SUP-XYZ'), tuple(meta_row)
    meta_import_host.destroy()
    # Legacy catalogue updates must not erase metadata added by newer versions.
    with connect() as c:
        c.execute("INSERT INTO products(name,sale_price_cents,active,sku,alias,supplier_code) VALUES('TEST Legacy Preserve',900,1,'LEG-SKU','Keep Alias','KEEP-SUP')")
        legacy_preserve_pid=c.execute("SELECT last_insert_rowid()").fetchone()[0]
        c.execute("INSERT INTO product_barcodes(product_id,barcode,qty_multiplier) VALUES(?,?,1)",(legacy_preserve_pid,'LEGPRES123'))
        c.execute("INSERT INTO quantity_prices(product_id,min_qty,unit_price_cents,pricing_mode) VALUES(?,?,?,?)",(legacy_preserve_pid,2,700,'UNIT'))
    legacy_update_path=Path(temp.name)/'catalogue-legacy-preserve.xlsx'
    wb=Workbook();ws=wb.active
    ws.append(["product key","barcode","article","famille","prix achat","prix vente","stock","alerte","barcode label","multiplicateur","prix pack","sku","fraction"])
    ws.append(["LEGACY-PRES","LEGPRES123","TEST Legacy Preserve","Général","0.00","9.00",0,0,"",1,None,"LEG-SKU",0]);wb.save(legacy_update_path);wb.close()
    legacy_update_host=__import__('tkinter').Toplevel(app)
    legacy_update=ProductsFrame(legacy_update_host);legacy_update.pack(fill='both',expand=True);app.update()
    def accept_legacy_update_preview():
        preview=next(w for w in legacy_update.winfo_children() if w.winfo_class()=='Toplevel')
        button(preview,'Importer 1').invoke()
    app.after(150,accept_legacy_update_preview)
    with patch('screens.products.filedialog.askopenfilename',return_value=str(legacy_update_path)), \
         patch('screens.products.messagebox.askyesnocancel',return_value=True):
        legacy_update.import_excel()
    with connect() as c:
        legacy_meta=c.execute("SELECT alias,supplier_code FROM products WHERE id=?",(legacy_preserve_pid,)).fetchone()
        legacy_offers=[tuple(r) for r in c.execute("SELECT min_qty,unit_price_cents,pricing_mode FROM quantity_prices WHERE product_id=?",(legacy_preserve_pid,)).fetchall()]
        assert tuple(legacy_meta)==('Keep Alias','KEEP-SUP'), tuple(legacy_meta)
        assert legacy_offers==[(2.0,700,'UNIT')], legacy_offers
    legacy_update_host.destroy()
    # Image-only products must carry their image through catalogue exchange and
    # receive a fresh device-local TODO-* barcode after import.
    from PIL import Image as PILImage
    image_source=Path(temp.name)/'exchange-photo.png'
    PILImage.new('RGB',(24,24),'white').save(image_source)
    with connect() as c:
        c.execute("INSERT INTO products(name,sale_price_cents,active,image_path) VALUES('TEST Image Exchange',600,1,'')")
        image_pid=c.execute("SELECT last_insert_rowid()").fetchone()[0]
    from services.images import import_image as import_product_image
    imported_rel=import_product_image(str(image_source))
    with connect() as c:
        c.execute("UPDATE products SET image_path=? WHERE id=?",(imported_rel,image_pid))
        c.execute("INSERT INTO product_barcodes(product_id,barcode,qty_multiplier) VALUES(?,?,1)",(image_pid,f'TODO-{image_pid:08d}'))
    image_export=Path(temp.name)/'catalogue-image-exchange.xlsx'
    image_export_host=__import__('tkinter').Toplevel(app)
    image_products=ProductsFrame(image_export_host);image_products.pack(fill='both',expand=True);app.update()
    with patch('screens.products.filedialog.asksaveasfilename',return_value=str(image_export)):
        image_products.export_catalogue()
    image_export_host.destroy()
    image_wb=load_workbook(image_export)
    image_ws=image_wb.active
    image_headers=[cell.value for cell in image_ws[1]]
    image_name_col=image_headers.index('article')+1
    for row_idx in range(image_ws.max_row,1,-1):
        if image_ws.cell(row_idx,image_name_col).value!='TEST Image Exchange':
            image_ws.delete_rows(row_idx,1)
    image_wb.save(image_export);image_wb.close()
    with connect() as c:
        c.execute("DELETE FROM products WHERE id=?",(image_pid,))
    image_import_host=__import__('tkinter').Toplevel(app)
    image_import=ProductsFrame(image_import_host);image_import.pack(fill='both',expand=True);app.update()
    def accept_image_preview():
        preview=next(w for w in image_import.winfo_children() if w.winfo_class()=='Toplevel')
        button(preview,'Importer 1').invoke()
    app.after(150,accept_image_preview)
    with patch('screens.products.filedialog.askopenfilename',return_value=str(image_export)):
        image_import.import_excel()
    with connect() as c:
        image_row=c.execute("SELECT id,image_path FROM products WHERE active=1 AND name='TEST Image Exchange'").fetchone()
        assert image_row and image_row['image_path'], image_row
        image_barcode=c.execute("SELECT barcode FROM product_barcodes WHERE product_id=?",(image_row['id'],)).fetchone()
        assert image_barcode and image_barcode['barcode']==f"TODO-{image_row['id']:08d}", image_barcode
    from services.images import abs_image as absolute_product_image
    assert absolute_product_image(image_row['image_path']).exists()
    image_import_host.destroy()
    sale.render_products();app.update()
    tactile_names=[w.cget('text') for w in descendants(sale.card_inner) if w.winfo_class()=='Label']
    assert 'TEST - Rice' not in tactile_names, 'Products without images must stay out of tactile grid'
    assert 'TEST Photo Internal' in tactile_names, 'Image-only products with TODO internal barcode must remain in tactile grid'
    assert 'TEST Photo Regular' not in tactile_names, 'Products with a real barcode must stay out of tactile grid'
    sale.query.set('TEST123');sale.handle_scan_input()
    assert sale.scan_submit_job is not None, 'Known barcode typing must schedule the 500 ms auto-submit'
    sale.confirm_search()
    assert sale.scan_submit_job is None, 'Scanner Enter must cancel pending auto-submit to prevent double lines'
    assert len(sale.cart)==1 and sale.cart[0]['qty']==1
    sale.change(1)
    # Invoke reference menu actions while preserving the current ticket.
    sale.functions();app.update()
    menu=next(w for w in sale.winfo_children() if w.winfo_class()=='Toplevel' and ('Fonctions' in str(w.title()) or 'الوظائف' in str(w.title())))
    for label in ['Duplicata','Modifier quantité','Modifier prix','Supprimer','Grille de prix','Clôture','Dépenses','Rapport','Raccourcis']:
        visible(button(menu,label))
    with patch('tkinter.simpledialog.askstring',return_value='4.00'):
        button(menu,'Modifier prix').invoke()
    assert sale.totals()[1]==800
    # End-to-end PRIX_1: grid price is selectable from the sale UI.
    with connect() as c:
        gid=c.execute("INSERT INTO price_grids(name,active) VALUES('TEST PRO',1)").lastrowid
        c.execute("INSERT INTO product_grid_prices(product_id,grid_id,unit_price_cents) VALUES(?,?,?)",(pid,gid,250))
    sale.restore_price()
    sale.cart[0].pop('manual_unit_price',None)
    sale.functions();app.update()
    menu=next(w for w in descendants(app) if w.winfo_class()=='Toplevel')
    button(menu,'Grille de prix').invoke();app.update()
    grid_dialog=next(w for w in descendants(app) if w.winfo_class()=='Toplevel')
    button(grid_dialog,'TEST PRO').invoke();app.update()
    assert sale.price_grid_id==gid
    assert sale.totals()[1]==500
    # Return to the normal grid so later acceptance scenarios are isolated.
    sale.price_grid_id=None;sale.price_grid_name=sale.tr('Normal','عادي')
    sale.cart[0].pop('manual_unit_price',None)
    sale.restore_price()
    assert sale.totals()[1]==600
    with patch('tkinter.simpledialog.askstring',return_value='Test held ticket'):
        sale.hold()
    assert not sale.cart
    sale.show_held();app.update()
    held=next(w for w in descendants(app) if w.winfo_class()=='Toplevel')
    tree=next(w for w in descendants(held) if w.winfo_class()=='Treeview')
    tree.selection_set(tree.get_children()[0]);button(held,'Reprendre').invoke()
    assert sale.cart[0]['qty']==2
    saved=list(sale.cart)
    app.lock_cashier();app.update()
    lock=app.lock_window
    app.show('products');assert app.current is sale
    lock.pin.set('9999');lock.unlock();assert app.lock_window is lock
    lock.pin.set('1234');lock.unlock();assert app.lock_window is None
    assert sale.cart==saved
    sale.calculator();app.update()
    calc=next(w for w in descendants(app) if w.winfo_class()=='Toplevel')
    calc.expression.set('12,5 * 2');button(calc,'=').invoke()
    assert calc.expression.get()=='25'
    button(calc,'Fermer').invoke()
    app.choose_theme();app.update()
    theme=next(w for w in descendants(app) if w.winfo_class()=='Toplevel')
    button(theme,'Vert').invoke()
    from database import get_setting
    assert get_setting('theme')=='Vert' and app.theme_color=='#15803D'
    assert sale.cart==saved
    sale.add_misc();app.update()
    misc=next(w for w in descendants(app) if hasattr(w,'misc_price_var'))
    misc.misc_name_var.set('TEST supplement')
    misc.misc_price_var.set('2.50')
    misc.misc_qty_var.set('1')
    misc.misc_confirm.invoke();app.update()
    assert sale.cart[-1]['is_misc'] and sale.totals()[1]==850
    sale.remove();assert sale.totals()[1]==600
    with patch('tkinter.simpledialog.askfloat',return_value=1):sale.discount()
    assert sale.totals()[1]==500
    for label in ['SOLDER avec','SOLDER sans','Fonctions']:visible(button(sale,label))
    def finish_payment():
        dialog=next(w for w in descendants(app) if isinstance(w,PaymentDialog))
        visible(dialog.confirm_button)
        dialog.amount.set('10.00')
        dialog.confirm_button.invoke()
    app.after(150,finish_payment)
    button(sale,'SOLDER avec').invoke();app.update()
    assert not sale.cart
    assert sale.last_sale_snapshot and sale.ticket.get_children()
    # Checkout stays inline: no automatic receipt popup. Cashier focus returns
    # to the barcode/search field and the compact blue reminder shows the result.
    assert not any(w for w in descendants(app) if w.winfo_class()=='Toplevel')
    assert sale.focus_get() is sale.entry
    reminder=sale.last_sale_summary.cget('text')
    assert 'PAYÉ' in reminder and 'RENDU' in reminder
    with connect() as c:
        row=dict(c.execute('SELECT * FROM sales').fetchone())
        assert row['total_cents']==500 and row['change_cents']==500
        assert c.execute('SELECT stock_qty FROM products WHERE id=?',(pid,)).fetchone()[0]==3
    # Receipt/PDF remains available on demand, but is never forced after checkout.
    sale.show_receipt(row);app.update()
    receipt=next(w for w in descendants(app) if w.winfo_class()=='Toplevel')
    visible(button(receipt,'PDF'))
    output=Path(os.environ.get('TODO_TEST_OUTPUT',str(Path(temp.name)/'receipt.pdf')))
    with patch('tkinter.filedialog.asksaveasfilename',return_value=str(output)):
        button(receipt,'PDF').invoke()
    assert output.read_bytes().startswith(b'%PDF-')
    receipt.destroy()
    # Logica-compatible behavior: completed ticket remains visible until the next article.
    sale.query.set('TEST123');sale.confirm_search();app.update()
    assert sale.cart and sale.last_sale_snapshot is None
    assert len(sale.ticket.get_children())==1
    sale.clear()
    sale.cash_tools();app.update()
    cash_window=next(w for w in descendants(app) if w.winfo_class()=='Toplevel')
    with patch('tkinter.simpledialog.askstring',return_value='TEST expense'), patch('tkinter.simpledialog.askfloat',return_value=1):
        button(cash_window,'Dépense').invoke()
    with connect() as c:
        assert c.execute('SELECT amount_cents FROM expenses').fetchone()[0]==100
    # Closing uses the real denomination counter instead of askfloat.
    def finish_cash_count():
        counter=next(w for w in descendants(app) if w.winfo_class()=='Toplevel' and 'Comptage' in str(w.title()))
        for value,count in ((100,1),(2,2)):
            for _ in range(count):button(counter,f'+ {value:g} DH').invoke()
        button(counter,'Valider le comptage').invoke()
    app.after(150,finish_cash_count)
    button(cash_window,'Clôturer').invoke()
    with connect() as c:
        closing=c.execute('SELECT * FROM cash_sessions').fetchone()
        assert closing['status']=='CLOSED' and closing['actual_cash_cents']==10400 and closing['difference_cents']==0
    cash_window.destroy()
    app.show('journal');app.update()
    # Complete a customer credit sale and settlement through actual Tk controls.
    from screens.clients import ClientEditor, PaymentsWindow
    from services.clients import get_client
    from services.cash import close_session
    open_session(app.user['id'],0)
    app.show('clients');app.update()
    button(app.current,'Nouveau').invoke();app.update()
    customer_editor=next(w for w in descendants(app) if isinstance(w,ClientEditor))
    customer_editor.name.set('TEST credit customer')
    button(customer_editor,'Enregistrer').invoke();app.update()
    with connect() as c:cid=c.execute("SELECT id FROM clients WHERE name='TEST credit customer'").fetchone()[0]
    app.show('sale');app.update();sale=app.current
    button(sale,'F6 Client').invoke();app.update()
    picker=next(w for w in descendants(app) if w.winfo_class()=='Toplevel')
    customer_tree=next(w for w in descendants(picker) if w.winfo_class()=='Treeview')
    customer_tree.selection_set(str(cid));button(picker,'Choisir').invoke()
    assert sale.client_id==cid
    sale.query.set('TEST123');sale.confirm_search()
    def credit_payment():
        dialog=next(w for w in descendants(app) if isinstance(w,PaymentDialog))
        dialog.choose_method('CREDIT');dialog.amount.set('1.00');app.update()
        visible(dialog.confirm_button)
        assert not dialog.confirm_button.instate(['disabled'])
        dialog.confirm_button.invoke()
    app.after(150,credit_payment)
    button(sale,'SOLDER sans').invoke();app.update()
    credit_summary=sale.last_sale_summary.cget('text')
    assert ('CRÉDIT' in credit_summary or 'دين' in credit_summary), 'Completed credit sale must preserve CREDIT in the last-sale summary'
    assert get_client(cid)['balance_cents']==200 and sale.client_id is None
    app.show('clients');app.update()
    app.current.tree.selection_set(str(cid));button(app.current,'Règlements').invoke();app.update()
    payments=next(w for w in descendants(app) if isinstance(w,PaymentsWindow))
    with patch('tkinter.simpledialog.askstring',side_effect=['2.00','CASH','TEST settlement']):
        button(payments,'Ajouter un règlement').invoke()
    assert get_client(cid)['balance_cents']==0
    payments.destroy()
    assert close_session(__import__('services.cash',fromlist=['get_open_session']).get_open_session()['id'],300)[:2]==(300,0)
    app.show('settings');app.update()
    settings=app.current
    with patch('services.printers.installed_printers',return_value=['Receipt Test']):
        button(settings,'Actualiser imprimantes').invoke()
    assert 'Receipt Test' in settings.printer_choice['values']
    settings.printer.set('Receipt Test');settings.print_mode.set('never')
    settings.drawer_enabled.set(True);settings.drawer_pin.set('1')
    button(settings,'Enregistrer impression').invoke()
    assert get_setting('printer_name')=='Receipt Test'
    assert get_setting('drawer_enabled')=='1' and get_setting('drawer_pin')=='1'
    # New Logica-compatible switches must persist through the real Settings UI save path.
    settings.credit_enabled.set(False)
    settings.closure_enabled.set(False)
    settings.sans_ticket_enabled.set(False)
    settings.payment_window_enabled.set(False)
    settings.customer_enabled.set(False)
    button(settings,'Enregistrer').invoke();app.update()
    assert get_setting('credit_enabled')=='0'
    assert get_setting('closure_enabled')=='0'
    assert get_setting('sans_ticket_enabled')=='0'
    assert get_setting('payment_window_enabled')=='0'
    assert get_setting('customer_display_enabled')=='0'
    print('PASS: UI product/category -> purchase confirmation -> stock 5 -> scan/qty/discount -> visible SOLDER/VALIDER at 1100x640 -> sale 5 DH -> stock 3 -> receipt PDF -> journal',flush=True)
    if os.environ.get('TODO_REVIEW_UI'):
        app.show('sale')
        app.title('ToDo POS - isolated acceptance test')
        app.mainloop()
    else:app.destroy()
temp.cleanup()
