import csv
import tkinter as tk
from xml.sax.saxutils import escape
from datetime import date, timedelta
from tkinter import ttk,filedialog,messagebox
from database import connect, get_setting
from services.money import fmt

class JournalFrame(ttk.Frame):
    def __init__(self,master):
        super().__init__(master,padding=10)
        self.lang=get_setting('language','fr');self.tr=lambda fr,ar: ar if self.lang=='ar' else fr
        top=ttk.Frame(self);top.pack(fill="x")
        ttk.Label(top,text=self.tr('Journal','السجل'),font=("Segoe UI",22,"bold")).pack(side="left")
        ttk.Button(top,text=self.tr('Export détaillé CSV','تصدير مفصل CSV'),command=self.export).pack(side="right");ttk.Button(top,text=self.tr('Export détaillé Excel','تصدير مفصل Excel'),command=self.export_excel).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Export détaillé PDF','تصدير مفصل PDF'),command=self.export_pdf).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Rapport articles','تقرير المنتجات'),command=self.article_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Rapport familles','تقرير الفئات'),command=self.family_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Rapport clients','تقرير الزبائن'),command=self.client_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Rapport vendeurs','تقرير البائعين'),command=self.seller_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Rapport caissiers','تقرير الكاشير'),command=self.cashier_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Rapport global','التقرير العام'),command=self.global_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Rapport jours','تقرير الأيام'),command=self.day_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Sans détails','بدون تفاصيل'),command=self.summary_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Export cumulé','تصدير تراكمي'),command=self.cumulative_export).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Paiements','الدفعات'),command=self.payment_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Retours','المرتجعات'),command=self.return_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Entête','الرؤوس'),command=self.header_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Journal détaillé','السجل المفصل'),command=self.movement_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Imprimer','طباعة'),command=self.print_report).pack(side="right",padx=5);ttk.Button(top,text=self.tr('Actualiser','تحديث'),command=self.refresh).pack(side="right",padx=5)
        filters=ttk.Frame(self);filters.pack(fill="x",pady=8)
        today=date.today();self.date_from=tk.StringVar(value=str(today));self.date_to=tk.StringVar(value=str(today));self.all_label=self.tr('Tous','الكل');self.cashier=tk.StringVar(value=self.all_label);self.seller=tk.StringVar(value=self.all_label);self.payment=tk.StringVar(value=self.all_label)
        for label,var,width in [(self.tr('Du','من'),self.date_from,11),(self.tr('Au','إلى'),self.date_to,11)]:ttk.Label(filters,text=label).pack(side="left");ttk.Entry(filters,textvariable=var,width=width).pack(side="left",padx=(3,10))
        ttk.Label(filters,text=self.tr('Caissier','الكاشير')).pack(side="left");self.cashier_box=ttk.Combobox(filters,textvariable=self.cashier,state="readonly",width=16);self.cashier_box.pack(side="left",padx=(3,10))
        ttk.Label(filters,text=self.tr('Vendeur','البائع')).pack(side="left");self.seller_box=ttk.Combobox(filters,textvariable=self.seller,state="readonly",width=16);self.seller_box.pack(side="left",padx=(3,10))
        ttk.Label(filters,text=self.tr('Paiement','الدفع')).pack(side="left");ttk.Combobox(filters,textvariable=self.payment,values=[self.all_label,"CASH","CARD","MIXED","CREDIT"],state="readonly",width=10).pack(side="left",padx=(3,10))
        ttk.Button(filters,text=self.tr('Aujourd’hui','اليوم'),command=lambda:self.set_period(0)).pack(side="left",padx=2);ttk.Button(filters,text=self.tr('7 jours','7 أيام'),command=lambda:self.set_period(6)).pack(side="left",padx=2);ttk.Button(filters,text=self.tr('30 jours','30 يوماً'),command=lambda:self.set_period(29)).pack(side="left",padx=2);ttk.Button(filters,text=self.tr('Consulter','عرض'),command=self.refresh).pack(side="right")
        cols=("id","ticket","date","cashier","seller","pay","total","cost","margin")
        self.t=ttk.Treeview(self,columns=cols,show="headings")
        for c,h,w in [("id","ID",45),("ticket",self.tr("Ticket","التذكرة"),190),("date",self.tr("Date","التاريخ"),160),("cashier",self.tr("Caissier","الكاشير"),110),("seller",self.tr("Vendeur","البائع"),110),("pay",self.tr("Paiement","الدفع"),90),("total",self.tr("Total","المجموع"),90),("cost",self.tr("Coût","التكلفة"),90),("margin",self.tr("Marge brute","الهامش الإجمالي"),100)]:self.t.heading(c,text=h);self.t.column(c,width=w,anchor="center")
        self.t.pack(fill="both",expand=True)
        self.summary=ttk.Label(self,text="",font=("Segoe UI",11,"bold"));self.summary.pack(anchor="e",pady=6)
        with connect() as c:
            names=[r[0] for r in c.execute("SELECT display_name FROM users WHERE active=1 ORDER BY display_name").fetchall()]
            sellers=[r[0] for r in c.execute("SELECT name FROM sellers ORDER BY name COLLATE NOCASE").fetchall()]
        self.cashier_box["values"]=[self.all_label,*names]
        self.seller_box["values"]=[self.all_label,*sellers]
        self.refresh()
    def set_period(self,days):
        end=date.today();self.date_to.set(str(end));self.date_from.set(str(end-timedelta(days=days)));self.refresh()
    def rows(self):
        try:
            date.fromisoformat(self.date_from.get());date.fromisoformat(self.date_to.get())
        except ValueError:
            messagebox.showerror(self.tr("Journal","السجل"),self.tr("Dates au format YYYY-MM-DD.","التواريخ يجب أن تكون بصيغة YYYY-MM-DD."),parent=self);return []
        where=["date(s.created_at)>=?","date(s.created_at)<=?"];params=[self.date_from.get(),self.date_to.get()]
        if self.cashier.get()!=self.all_label:where.append("u.display_name=?");params.append(self.cashier.get())
        if self.seller.get()!=self.all_label:where.append("v.name=?");params.append(self.seller.get())
        if self.payment.get()!=self.all_label:
            if self.payment.get() in ("CASH","CARD"):
                where.append("EXISTS (SELECT 1 FROM sale_payments sp WHERE sp.sale_id=s.id AND sp.payment_method=?)");params.append(self.payment.get())
            else:
                where.append("s.payment_method=?");params.append(self.payment.get())
        sql="""SELECT s.id,s.sale_no,s.created_at,u.display_name,COALESCE(v.name,'') seller_name,s.payment_method,
          COALESCE((SELECT SUM(sp.amount_cents) FROM sale_payments sp WHERE sp.sale_id=s.id AND sp.payment_method=\'CASH\'),0) cash_paid,
          COALESCE((SELECT SUM(sp.amount_cents) FROM sale_payments sp WHERE sp.sale_id=s.id AND sp.payment_method=\'CARD\'),0) card_paid,
          COALESCE((SELECT SUM(rp.amount_cents) FROM return_payments rp JOIN returns rr ON rr.id=rp.return_id WHERE rr.sale_id=s.id AND rp.payment_method=\'CASH\'),0) cash_refund,
          COALESCE((SELECT SUM(rp.amount_cents) FROM return_payments rp JOIN returns rr ON rr.id=rp.return_id WHERE rr.sale_id=s.id AND rp.payment_method=\'CARD\'),0) card_refund,
          s.total_cents-COALESCE((SELECT SUM(r.total_cents) FROM returns r WHERE r.sale_id=s.id),0) total_cents,
          COALESCE((SELECT SUM(si.cost_price_cents*si.qty) FROM sale_items si WHERE si.sale_id=s.id),0)
          -COALESCE((SELECT SUM(ri.qty*si.cost_price_cents) FROM return_items ri JOIN sale_items si ON si.id=ri.sale_item_id WHERE si.sale_id=s.id),0) cost
          FROM sales s JOIN users u ON u.id=s.cashier_user_id LEFT JOIN sellers v ON v.id=s.seller_id WHERE """+" AND ".join(where)+" ORDER BY s.id DESC LIMIT 5000"
        with connect() as c:return c.execute(sql,params).fetchall()
    def refresh(self):
        rows=self.rows();self.t.delete(*self.t.get_children());total=cost=0
        for r in rows:
            margin=int(r["total_cents"]-r["cost"]);total+=r["total_cents"];cost+=r["cost"]
            pay=r["payment_method"]
            if pay=="MIXED":pay="MIXED (Cash {} + Card {})".format(fmt(r["cash_paid"]-r["cash_refund"],""),fmt(r["card_paid"]-r["card_refund"],""))
            self.t.insert("", "end",values=(r["id"],r["sale_no"],r["created_at"],r["display_name"],r["seller_name"],pay,fmt(r["total_cents"],""),fmt(r["cost"],""),fmt(margin,"")))
        self.summary.config(text=f"{len(rows)} ticket(s) · Ventes nettes {fmt(total)} · Coût {fmt(cost)} · Marge brute {fmt(total-cost)}")
    def detail_report(self,mode):
        try:
            date.fromisoformat(self.date_from.get());date.fromisoformat(self.date_to.get())
        except ValueError:
            messagebox.showerror(self.tr("Journal","السجل"),self.tr("Dates au format YYYY-MM-DD.","التواريخ يجب أن تكون بصيغة YYYY-MM-DD."),parent=self);return
        if mode=="article":
            group="p.id,p.name";label="Article";select="p.name"
        else:
            group="COALESCE(cat.id,0),COALESCE(cat.name,'Sans famille')";label="Famille";select="COALESCE(cat.name,'Sans famille')"
        sql=f"""SELECT {select} label,
            SUM(si.qty-COALESCE((SELECT SUM(ri.qty) FROM return_items ri WHERE ri.sale_item_id=si.id),0)) qty,
            SUM(si.line_total_cents-COALESCE((SELECT SUM(ri.total_cents) FROM return_items ri WHERE ri.sale_item_id=si.id),0)) gross,
            SUM(si.cost_price_cents*(si.qty-COALESCE((SELECT SUM(ri.qty) FROM return_items ri WHERE ri.sale_item_id=si.id),0))) cost
            FROM sale_items si JOIN sales s ON s.id=si.sale_id JOIN products p ON p.id=si.product_id
            LEFT JOIN categories cat ON cat.id=p.category_id
            WHERE date(s.created_at)>=? AND date(s.created_at)<=? GROUP BY {group} ORDER BY gross DESC"""
        with connect() as c:rows=c.execute(sql,(self.date_from.get(),self.date_to.get())).fetchall()
        w=tk.Toplevel(self);w.title(self.tr(f"Rapport par {label}",f"تقرير حسب {label}"));w.geometry("850x560");w.transient(self.winfo_toplevel())
        tree=ttk.Treeview(w,columns=("label","qty","sales","cost","margin"),show="headings")
        for key,title,width in [("label",label,280),("qty","Qté",90),("sales","Ventes",120),("cost","Coût",120),("margin","Marge brute",120)]:tree.heading(key,text=title);tree.column(key,width=width,anchor="e" if key!="label" else "w")
        tree.pack(fill="both",expand=True,padx=12,pady=12)
        total=cost=0
        for r in rows:
            total+=r["gross"];cost+=r["cost"];tree.insert("","end",values=(r["label"],f"{r['qty']:g}",fmt(r["gross"],""),fmt(r["cost"],""),fmt(r["gross"]-r["cost"],"")))
        ttk.Label(w,text=f"Total {fmt(total)} · Coût {fmt(cost)} · Marge {fmt(total-cost)}",font=("Segoe UI",11,"bold")).pack(anchor="e",padx=12,pady=(0,12))
    def article_report(self):self.detail_report("article")
    def family_report(self):self.detail_report("family")
    def global_report(self):
        rows=self.rows();total=sum(r["total_cents"] for r in rows);cost=sum(r["cost"] for r in rows)
        cash=sum(r["cash_paid"]-r["cash_refund"] for r in rows);card=sum(r["card_paid"]-r["card_refund"] for r in rows)
        self._simple_report(self.tr("Rapport global","التقرير العام"),
            (self.tr("Indicateur","المؤشر"),self.tr("Valeur","القيمة")),
            [(self.tr("Période","الفترة"),f"{self.date_from.get()} → {self.date_to.get()}"),
             (self.tr("Tickets","التذاكر"),len(rows)),(self.tr("Ventes nettes","صافي المبيعات"),fmt(total,"")),
             (self.tr("Coût","التكلفة"),fmt(cost,"")),(self.tr("Marge brute","الهامش الإجمالي"),fmt(total-cost,"")),
             (self.tr("Cash net","صافي النقد"),fmt(cash,"")),(self.tr("Carte nette","صافي البطاقة"),fmt(card,""))])
    def cashier_report(self):
        try:
            date.fromisoformat(self.date_from.get());date.fromisoformat(self.date_to.get())
        except ValueError:
            messagebox.showerror(self.tr("Journal","السجل"),self.tr("Dates au format YYYY-MM-DD.","التواريخ يجب أن تكون بصيغة YYYY-MM-DD."),parent=self);return
        sql="""SELECT u.display_name label,COUNT(*) tickets,
            SUM(s.total_cents-COALESCE((SELECT SUM(r.total_cents) FROM returns r WHERE r.sale_id=s.id),0)) sales
            FROM sales s JOIN users u ON u.id=s.cashier_user_id
            WHERE s.status='COMPLETED' AND date(s.created_at)>=? AND date(s.created_at)<=?
            GROUP BY s.cashier_user_id ORDER BY sales DESC"""
        with connect() as c:rows=c.execute(sql,(self.date_from.get(),self.date_to.get())).fetchall()
        self._simple_report(self.tr("Rapport caissiers","تقرير الكاشير"),
            (self.tr("Caissier","الكاشير"),self.tr("Tickets","التذاكر"),self.tr("Ventes nettes","صافي المبيعات")),
            [(r["label"],r["tickets"],fmt(r["sales"] or 0,"")) for r in rows])

    def seller_report(self):
        try:
            date.fromisoformat(self.date_from.get());date.fromisoformat(self.date_to.get())
        except ValueError:
            messagebox.showerror(self.tr("Journal","السجل"),self.tr("Dates au format YYYY-MM-DD.","التواريخ يجب أن تكون بصيغة YYYY-MM-DD."),parent=self);return
        sql="""SELECT COALESCE(v.name,'Sans vendeur') label,COUNT(*) tickets,
            SUM(s.total_cents-COALESCE((SELECT SUM(r.total_cents) FROM returns r WHERE r.sale_id=s.id),0)) sales
            FROM sales s LEFT JOIN sellers v ON v.id=s.seller_id
            WHERE s.status='COMPLETED' AND date(s.created_at)>=? AND date(s.created_at)<=?
            GROUP BY s.seller_id ORDER BY sales DESC"""
        with connect() as c:rows=c.execute(sql,(self.date_from.get(),self.date_to.get())).fetchall()
        w=tk.Toplevel(self);w.title(self.tr("Rapport vendeurs","تقرير البائعين"));w.geometry("700x540");w.transient(self.winfo_toplevel())
        tree=ttk.Treeview(w,columns=("seller","tickets","sales"),show="headings")
        for key,title,width in [("seller",self.tr("Vendeur","البائع"),330),("tickets",self.tr("Tickets","التذاكر"),100),("sales",self.tr("Ventes nettes","صافي المبيعات"),150)]:tree.heading(key,text=title);tree.column(key,width=width,anchor="e" if key!="seller" else "w")
        tree.pack(fill="both",expand=True,padx=12,pady=12)
        total=0
        for r in rows:
            total+=r["sales"] or 0;tree.insert("","end",values=(r["label"],r["tickets"],fmt(r["sales"] or 0,"")))
        ttk.Label(w,text=self.tr("Total ventes nettes ","إجمالي صافي المبيعات ")+fmt(total),font=("Segoe UI",11,"bold")).pack(anchor="e",padx=12,pady=(0,12))

    def client_report(self):
        try:
            date.fromisoformat(self.date_from.get());date.fromisoformat(self.date_to.get())
        except ValueError:
            messagebox.showerror(self.tr("Journal","السجل"),self.tr("Dates au format YYYY-MM-DD.","التواريخ يجب أن تكون بصيغة YYYY-MM-DD."),parent=self);return
        sql="""SELECT COALESCE(cl.name,'Client comptoir') label,COUNT(DISTINCT s.id) tickets,
            SUM(s.total_cents-COALESCE((SELECT SUM(r.total_cents) FROM returns r WHERE r.sale_id=s.id),0)) sales
            FROM sales s LEFT JOIN clients cl ON cl.id=s.client_id
            WHERE s.status='COMPLETED' AND date(s.created_at)>=? AND date(s.created_at)<=?
            GROUP BY s.client_id ORDER BY sales DESC"""
        with connect() as c:rows=c.execute(sql,(self.date_from.get(),self.date_to.get())).fetchall()
        w=tk.Toplevel(self);w.title(self.tr("Rapport clients","تقرير الزبائن"));w.geometry("700x540");w.transient(self.winfo_toplevel())
        tree=ttk.Treeview(w,columns=("client","tickets","sales"),show="headings")
        for key,title,width in [("client",self.tr("Client","الزبون"),330),("tickets",self.tr("Tickets","التذاكر"),100),("sales",self.tr("Ventes nettes","صافي المبيعات"),150)]:tree.heading(key,text=title);tree.column(key,width=width,anchor="e" if key!="client" else "w")
        tree.pack(fill="both",expand=True,padx=12,pady=12)
        total=0
        for r in rows:
            total+=r["sales"] or 0;tree.insert("","end",values=(r["label"],r["tickets"],fmt(r["sales"] or 0,"")))
        ttk.Label(w,text=self.tr("Total ventes nettes ","إجمالي صافي المبيعات ")+fmt(total),font=("Segoe UI",11,"bold")).pack(anchor="e",padx=12,pady=(0,12))
    def day_report(self):
        try:
            date.fromisoformat(self.date_from.get());date.fromisoformat(self.date_to.get())
        except ValueError:
            messagebox.showerror(self.tr("Journal","السجل"),self.tr("Dates au format YYYY-MM-DD.","التواريخ يجب أن تكون بصيغة YYYY-MM-DD."),parent=self);return
        sql="""SELECT date(s.created_at,'localtime') day,COUNT(*) tickets,
            SUM(s.total_cents-COALESCE((SELECT SUM(r.total_cents) FROM returns r WHERE r.sale_id=s.id),0)) sales
            FROM sales s WHERE s.status='COMPLETED' AND date(s.created_at)>=? AND date(s.created_at)<=?
            GROUP BY day ORDER BY day"""
        with connect() as c:rows=c.execute(sql,(self.date_from.get(),self.date_to.get())).fetchall()
        w=tk.Toplevel(self);w.title(self.tr("Rapport par jours","تقرير حسب الأيام"));w.geometry("650x540");w.transient(self.winfo_toplevel())
        tree=ttk.Treeview(w,columns=("day","tickets","sales"),show="headings")
        for key,title,width in [("day",self.tr("Jour","اليوم"),180),("tickets",self.tr("Tickets","التذاكر"),100),("sales",self.tr("Ventes nettes","صافي المبيعات"),180)]:tree.heading(key,text=title);tree.column(key,width=width,anchor="center" if key!="sales" else "e")
        tree.pack(fill="both",expand=True,padx=12,pady=12)
        total=0
        for r in rows:
            total+=r["sales"] or 0;tree.insert("","end",values=(r["day"],r["tickets"],fmt(r["sales"] or 0,"")))
        ttk.Label(w,text=f"Total ventes nettes {fmt(total)}",font=("Segoe UI",11,"bold")).pack(anchor="e",padx=12,pady=(0,12))

    def print_report(self):
        rows=self.rows()
        if not rows:
            messagebox.showinfo("ToDo",self.tr("Aucune donnée à imprimer.","لا توجد بيانات للطباعة."),parent=self);return
        p=filedialog.asksaveasfilename(defaultextension=".pdf",filetypes=[("PDF","*.pdf")],title=self.tr("Imprimer le rapport","طباعة التقرير"))
        if not p:return
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4,landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer
            data=[[self.tr("Ticket","التذكرة"),self.tr("Date","التاريخ"),self.tr("Caissier","الكاشير"),self.tr("Vendeur","البائع"),self.tr("Paiement","الدفع"),self.tr("Total","المجموع")]]
            for r in rows:data.append([r["sale_no"],r["created_at"],r["display_name"],r["seller_name"],r["payment_method"],fmt(r["total_cents"],"")])
            doc=SimpleDocTemplate(p,pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
            table=Table(data,repeatRows=1,colWidths=[125,125,110,110,90,90]);table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#2563EB")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),0.3,colors.grey),("FONTSIZE",(0,0),(-1,-1),8)]))
            styles=getSampleStyleSheet();doc.build([Paragraph(self.tr("Rapport Journal","تقرير السجل"),styles["Title"]),Spacer(1,8),Paragraph(f"{self.date_from.get()} - {self.date_to.get()}",styles["Heading3"]),Spacer(1,8),table])
            import os
            if os.name=="nt":os.startfile(p,"print")
            else:messagebox.showinfo("ToDo",self.tr("PDF prêt pour impression.","ملف PDF جاهز للطباعة."),parent=self)
        except Exception as e:messagebox.showerror("ToDo",str(e),parent=self)

    def movement_report(self):
        try:
            date.fromisoformat(self.date_from.get());date.fromisoformat(self.date_to.get())
        except ValueError:
            messagebox.showerror(self.tr("Journal","السجل"),self.tr("Dates au format YYYY-MM-DD.","التواريخ يجب أن تكون بصيغة YYYY-MM-DD."),parent=self);return
        sql="""SELECT date(s.created_at) day,time(s.created_at) hour,si.name_snapshot article,
            si.unit_price_cents,si.qty,si.discount_cents,si.net_total_cents,
            u.display_name cashier,COALESCE(v.name,'') seller,COALESCE(cl.name,'Client comptoir') client
            FROM sale_items si JOIN sales s ON s.id=si.sale_id JOIN users u ON u.id=s.cashier_user_id
            LEFT JOIN sellers v ON v.id=s.seller_id LEFT JOIN clients cl ON cl.id=s.client_id
            WHERE s.status='COMPLETED' AND date(s.created_at)>=? AND date(s.created_at)<=?
            ORDER BY s.id DESC,si.id"""
        with connect() as c:rows=c.execute(sql,(self.date_from.get(),self.date_to.get())).fetchall()
        self._simple_report(self.tr("Journal détaillé","السجل المفصل"),
            (self.tr("Date","التاريخ"),self.tr("Heure","الوقت"),self.tr("Article","المنتوج"),self.tr("Prix V","ثمن البيع"),self.tr("Qté","الكمية"),self.tr("Remise","التخفيض"),self.tr("Total Net","الصافي"),self.tr("Caissier","الكاشير"),self.tr("Vendeur","البائع"),self.tr("Client","الزبون"),self.tr("Opération","العملية")),
            [(r["day"],r["hour"],r["article"],fmt(r["unit_price_cents"],""),f'{r["qty"]:g}',fmt(r["discount_cents"],""),fmt(r["net_total_cents"],""),r["cashier"],r["seller"],r["client"],self.tr("VENTE","بيع")) for r in rows])

    def header_report(self):
        try:
            date.fromisoformat(self.date_from.get());date.fromisoformat(self.date_to.get())
        except ValueError:
            messagebox.showerror(self.tr("Journal","السجل"),self.tr("Dates au format YYYY-MM-DD.","التواريخ يجب أن تكون بصيغة YYYY-MM-DD."),parent=self);return
        sql="""SELECT s.sale_no,s.created_at,u.display_name cashier,COALESCE(v.name,'') seller,
            COALESCE(cl.name,'Client comptoir') client,s.payment_method,s.total_cents
            FROM sales s JOIN users u ON u.id=s.cashier_user_id
            LEFT JOIN sellers v ON v.id=s.seller_id LEFT JOIN clients cl ON cl.id=s.client_id
            WHERE s.status='COMPLETED' AND date(s.created_at)>=? AND date(s.created_at)<=?
            ORDER BY s.id DESC"""
        with connect() as c:rows=c.execute(sql,(self.date_from.get(),self.date_to.get())).fetchall()
        self._simple_report(self.tr("Entête journal","رؤوس السجل"),
            (self.tr("Ticket","التذكرة"),self.tr("Date","التاريخ"),self.tr("Caissier","الكاشير"),self.tr("Vendeur","البائع"),self.tr("Client","الزبون"),self.tr("Paiement","الدفع"),self.tr("Total","المجموع")),
            [(r["sale_no"],r["created_at"],r["cashier"],r["seller"],r["client"],r["payment_method"],fmt(r["total_cents"],"")) for r in rows])

    def summary_report(self):
        rows=self.rows()
        total=sum(r["total_cents"] for r in rows);cost=sum(r["cost"] for r in rows)
        cash=sum((r["cash_paid"]-r["cash_refund"]) for r in rows)
        card=sum((r["card_paid"]-r["card_refund"]) for r in rows)
        self._simple_report(self.tr("Journal sans détails","السجل بدون تفاصيل"),(self.tr("Indicateur","المؤشر"),self.tr("Valeur","القيمة")),[
            (self.tr("Période","الفترة"),f"{self.date_from.get()} → {self.date_to.get()}"),
            (self.tr("Tickets","التذاكر"),str(len(rows))),(self.tr("Ventes nettes","صافي المبيعات"),fmt(total,"")),(self.tr("Coût","التكلفة"),fmt(cost,"")),
            (self.tr("Marge brute","الهامش الإجمالي"),fmt(total-cost,"")),(self.tr("Cash net","صافي النقد"),fmt(cash,"")),(self.tr("Carte nette","صافي البطاقة"),fmt(card,""))])

    def cumulative_export(self):
        rows=self.rows()
        total=sum(r["total_cents"] for r in rows);cost=sum(r["cost"] for r in rows)
        cash=sum((r["cash_paid"]-r["cash_refund"]) for r in rows);card=sum((r["card_paid"]-r["card_refund"]) for r in rows)
        p=filedialog.asksaveasfilename(defaultextension=".csv",filetypes=[("CSV","*.csv")],title=self.tr("Exporter cumul","تصدير التراكمي"))
        if not p:return
        labels=[self.tr("Période","الفترة"),self.tr("Tickets","التذاكر"),self.tr("Ventes nettes","صافي المبيعات"),self.tr("Coût","التكلفة"),self.tr("Marge brute","الهامش الإجمالي"),self.tr("Cash net","صافي النقد"),self.tr("Carte nette","صافي البطاقة")]
        vals=[f"{self.date_from.get()} -> {self.date_to.get()}",len(rows),total,cost,total-cost,cash,card]
        with open(p,"w",newline="",encoding="utf-8-sig") as fh:
            wr=csv.writer(fh);wr.writerow([self.tr("Indicateur","المؤشر"),self.tr("Valeur","القيمة")])
            wr.writerows(zip(labels,vals))
        messagebox.showinfo("ToDo",self.tr("Export cumulé terminé.","تم التصدير التراكمي."),parent=self)

    def payment_report(self):
        try:
            date.fromisoformat(self.date_from.get());date.fromisoformat(self.date_to.get())
        except ValueError:
            messagebox.showerror(self.tr("Journal","السجل"),self.tr("Dates au format YYYY-MM-DD.","التواريخ يجب أن تكون بصيغة YYYY-MM-DD."),parent=self);return
        sql="""SELECT sp.payment_method method,SUM(sp.amount_cents) paid,
            COALESCE((SELECT SUM(rp.amount_cents) FROM return_payments rp JOIN returns r ON r.id=rp.return_id
              WHERE rp.payment_method=sp.payment_method AND date(r.created_at)>=? AND date(r.created_at)<=?),0) refunded
            FROM sale_payments sp JOIN sales s ON s.id=sp.sale_id
            WHERE s.status='COMPLETED' AND date(s.created_at)>=? AND date(s.created_at)<=? GROUP BY sp.payment_method ORDER BY sp.payment_method"""
        with connect() as c:rows=c.execute(sql,(self.date_from.get(),self.date_to.get(),self.date_from.get(),self.date_to.get())).fetchall()
        self._simple_report(self.tr("Paiements","الدفعات"),(self.tr("Mode","الطريقة"),self.tr("Encaissé","المقبوض"),self.tr("Remboursé","المسترجع"),self.tr("Net","الصافي")),[(r["method"],fmt(r["paid"] or 0,""),fmt(r["refunded"] or 0,""),fmt((r["paid"] or 0)-(r["refunded"] or 0),"")) for r in rows])

    def return_report(self):
        try:
            date.fromisoformat(self.date_from.get());date.fromisoformat(self.date_to.get())
        except ValueError:
            messagebox.showerror(self.tr("Journal","السجل"),self.tr("Dates au format YYYY-MM-DD.","التواريخ يجب أن تكون بصيغة YYYY-MM-DD."),parent=self);return
        sql="""SELECT r.return_no,r.created_at,s.sale_no,r.refund_method,r.total_cents
            FROM returns r JOIN sales s ON s.id=r.sale_id WHERE date(r.created_at)>=? AND date(r.created_at)<=? ORDER BY r.id DESC"""
        with connect() as c:rows=c.execute(sql,(self.date_from.get(),self.date_to.get())).fetchall()
        self._simple_report(self.tr("Retours","المرتجعات"),(self.tr("Retour","المرتجع"),self.tr("Date","التاريخ"),self.tr("Ticket","التذكرة"),self.tr("Mode","الطريقة"),self.tr("Montant","المبلغ")),[(r["return_no"],r["created_at"],r["sale_no"],r["refund_method"],fmt(r["total_cents"],"")) for r in rows])

    def _simple_report(self,title,headers,rows):
        w=tk.Toplevel(self);w.title(title);w.geometry("820x540");w.transient(self.winfo_toplevel())
        cols=tuple(f"c{i}" for i in range(len(headers)));tree=ttk.Treeview(w,columns=cols,show="headings")
        for i,(c,h) in enumerate(zip(cols,headers)):tree.heading(c,text=h);tree.column(c,width=150 if i else 180,anchor="center")
        tree.pack(fill="both",expand=True,padx=12,pady=12)
        for row in rows:tree.insert("","end",values=row)

    def export_excel(self):
        p=filedialog.asksaveasfilename(defaultextension=".xlsx",filetypes=[("Excel","*.xlsx")],title=self.tr("Exporter journal Excel","تصدير سجل المبيعات Excel"))
        if not p:return
        rows=self.rows()
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font,PatternFill,Alignment
            wb=Workbook();ws=wb.active;ws.title=self.tr("Journal ventes","سجل المبيعات")
            headers=["ID","Ticket","Date","Caissier","Vendeur","Paiement","Total","Coût","Marge brute"];ws.append(headers)
            for cell in ws[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="2563EB");cell.alignment=Alignment(horizontal="center")
            total=cost=0
            for r in rows:
                total+=r["total_cents"];cost+=r["cost"];pay=r["payment_method"]
                if pay=="MIXED":pay="MIXED (Cash {} + Card {})".format(fmt(r["cash_paid"]-r["cash_refund"],""),fmt(r["card_paid"]-r["card_refund"],""))
                ws.append([r["id"],r["sale_no"],r["created_at"],r["display_name"],r["seller_name"],pay,r["total_cents"]/100,r["cost"]/100,(r["total_cents"]-r["cost"])/100])
            ws.append([]);ws.append(["","","","","","TOTAL",total/100,cost/100,(total-cost)/100])
            for col,width in {"A":8,"B":22,"C":20,"D":18,"E":18,"F":14,"G":14,"H":14,"I":16}.items():ws.column_dimensions[col].width=width
            for row in ws.iter_rows(min_row=2,min_col=7,max_col=9):
                for cell in row:cell.number_format='#,##0.00'
            ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions;wb.save(p);messagebox.showinfo("ToDo",self.tr("Excel exporté.","تم تصدير Excel."),parent=self)
        except Exception as e:messagebox.showerror("ToDo",str(e),parent=self)
    def export_pdf(self):
        p=filedialog.asksaveasfilename(defaultextension=".pdf",filetypes=[("PDF","*.pdf")],title=self.tr("Exporter journal PDF","تصدير سجل المبيعات PDF"))
        if not p:return
        rows=self.rows()
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer
            styles=getSampleStyleSheet();doc=SimpleDocTemplate(p,pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
            data=[["Ticket","Date","Caissier","Vendeur","Paiement","Total","Coût","Marge"]];total=cost=0
            for r in rows:
                total+=r["total_cents"];cost+=r["cost"];pay=r["payment_method"]
                if pay=="MIXED":pay="MIXED (Cash {} + Card {})".format(fmt(r["cash_paid"]-r["cash_refund"],""),fmt(r["card_paid"]-r["card_refund"],""))
                data.append([r["sale_no"],r["created_at"],r["display_name"],r["seller_name"],pay,fmt(r["total_cents"],""),fmt(r["cost"],""),fmt(r["total_cents"]-r["cost"],"")])
            table=Table(data,repeatRows=1,colWidths=[105,105,80,80,65,70,70,70]);table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#2563EB")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),0.3,colors.grey),("FONTSIZE",(0,0),(-1,-1),8),("ALIGN",(5,1),(-1,-1),"RIGHT"),("VALIGN",(0,0),(-1,-1),"MIDDLE")]))
            title=Paragraph(f"Journal des ventes — {escape(self.date_from.get())} au {escape(self.date_to.get())}",styles["Title"])
            summary=Paragraph(f"{len(rows)} ticket(s) — Ventes nettes {fmt(total)} — Coût {fmt(cost)} — Marge brute {fmt(total-cost)}",styles["Heading3"])
            doc.build([title,Spacer(1,10),summary,Spacer(1,10),table]);messagebox.showinfo("ToDo",self.tr("PDF exporté.","تم تصدير PDF."),parent=self)
        except Exception as e:messagebox.showerror("ToDo",str(e),parent=self)
    def export(self):
        p=filedialog.asksaveasfilename(defaultextension=".csv",filetypes=[("CSV","*.csv")],title=self.tr("Exporter journal","تصدير السجل"))
        if not p:return
        rows=self.rows()
        with open(p,"w",newline="",encoding="utf-8-sig") as f:
            wr=csv.writer(f);wr.writerow(["ID",self.tr("Ticket","التذكرة"),self.tr("Date","التاريخ"),self.tr("Caissier","الكاشير"),self.tr("Vendeur","البائع"),self.tr("Paiement","الدفع"),self.tr("Total cents","المجموع بالسنتيم"),self.tr("Cost cents","التكلفة بالسنتيم"),self.tr("Margin cents","الهامش بالسنتيم")])
            for r in rows:
                pay=r["payment_method"]
                if pay=="MIXED":pay="MIXED (Cash {} + Card {})".format(fmt(r["cash_paid"]-r["cash_refund"],""),fmt(r["card_paid"]-r["card_refund"],""))
                wr.writerow([r["id"],r["sale_no"],r["created_at"],r["display_name"],r["seller_name"],pay,r["total_cents"],r["cost"],r["total_cents"]-r["cost"]])
        messagebox.showinfo("ToDo",self.tr("Export terminé.","تم التصدير."),parent=self)
